import io
import os
import zipfile
from collections.abc import Iterator
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import docx
from pypdf import PdfReader

try:
    from pptx import Presentation
except ImportError:  # python-pptx is optional; PPTX just won't be supported without it
    Presentation = None


@dataclass(frozen=True)
class ExtractedSegment:
    text: str
    page: int | None = None
    section: str | None = None


def normalize_file_type(file_type: str, path: str | Path | None = None) -> str:
    value = (file_type or "").lower().split(";", 1)[0].strip()
    suffix = Path(path).suffix.lower() if path else ""
    if value in {"application/pdf", "pdf"} or suffix == ".pdf":
        return "pdf"
    if value in {"application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"} or suffix == ".docx":
        return "docx"
    if value in {"application/vnd.openxmlformats-officedocument.presentationml.presentation", "pptx"} or suffix == ".pptx":
        return "pptx"
    if value in {"application/zip", "application/x-zip-compressed", "zip"} or suffix == ".zip":
        return "zip"
    return "text"


def extract_metadata(path: str | Path, file_type: str) -> dict[str, Any]:
    file_path = Path(path)
    kind = normalize_file_type(file_type, file_path)
    metadata: dict[str, Any] = {
        "title": file_path.stem,
        "author": None,
        "page_count": None,
        "file_name": file_path.name,
        "file_size": file_path.stat().st_size,
        "format": kind,
    }

    if kind == "pdf":
        reader = PdfReader(str(file_path))
        pdf_meta = reader.metadata or {}
        metadata.update({
            "title": _clean_metadata_value(getattr(pdf_meta, "title", None)) or file_path.stem,
            "author": _clean_metadata_value(getattr(pdf_meta, "author", None)),
            "subject": _clean_metadata_value(getattr(pdf_meta, "subject", None)),
            "creator": _clean_metadata_value(getattr(pdf_meta, "creator", None)),
            "producer": _clean_metadata_value(getattr(pdf_meta, "producer", None)),
            "creation_date": str(getattr(pdf_meta, "creation_date", "") or "") or None,
            "page_count": len(reader.pages),
        })
    elif kind == "docx":
        document = docx.Document(str(file_path))
        props = document.core_properties
        metadata.update({
            "title": props.title or file_path.stem,
            "author": props.author or None,
            "subject": props.subject or None,
            "keywords": props.keywords or None,
            "created": props.created.isoformat() if props.created else None,
            "modified": props.modified.isoformat() if props.modified else None,
            "section_count": len(document.sections),
        })
    elif kind == "pptx" and Presentation is not None:
        prs = Presentation(str(file_path))
        props = prs.core_properties
        metadata.update({
            "title": props.title or file_path.stem,
            "author": props.author or None,
            "subject": props.subject or None,
            "keywords": props.keywords or None,
            "created": props.created.isoformat() if props.created else None,
            "modified": props.modified.isoformat() if props.modified else None,
            "page_count": len(prs.slides),
        })
    return {key: value for key, value in metadata.items() if value is not None}


def iter_extracted_segments(path: str | Path, file_type: str) -> Iterator[ExtractedSegment]:
    file_path = Path(path)
    kind = normalize_file_type(file_type, file_path)
    if kind == "pdf":
        yield from _iter_pdf(file_path)
    elif kind == "docx":
        yield from _iter_docx(file_path)
    elif kind == "pptx":
        yield from _iter_pptx(file_path)
    elif kind == "zip":
        yield from _iter_zip(file_path)
    else:
        yield from _iter_text(file_path)


def _iter_pdf(path: Path) -> Iterator[ExtractedSegment]:
    reader = PdfReader(str(path))
    for page_number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            yield ExtractedSegment(text=text, page=page_number, section=f"Page {page_number}")


def _iter_docx(path: Path) -> Iterator[ExtractedSegment]:
    document = docx.Document(str(path))
    current_section: str | None = None
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = (paragraph.style.name if paragraph.style else "").lower()
        if style_name.startswith("heading"):
            current_section = text
        yield ExtractedSegment(text=text, section=current_section)

    for table_index, table in enumerate(document.tables, start=1):
        for row in table.rows:
            text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if text:
                yield ExtractedSegment(text=text, section=f"Table {table_index}")


def _iter_pptx(path: Path) -> Iterator[ExtractedSegment]:
    if Presentation is None:
        raise ValueError("python-pptx is not installed; cannot read .pptx files")
    prs = Presentation(str(path))
    for slide_number, slide in enumerate(prs.slides, start=1):
        lines: list[str] = []
        title = None
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for para in shape.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs).strip() or para.text.strip()
                if text:
                    lines.append(text)
            if title is None and getattr(shape, "text", "").strip():
                title = shape.text.strip().splitlines()[0][:80]
        # Include speaker notes — often the richest explanation.
        try:
            if slide.has_notes_slide:
                notes = (slide.notes_slide.notes_text_frame.text or "").strip()
                if notes:
                    lines.append(f"[Notes] {notes}")
        except Exception:
            pass
        body = "\n".join(lines).strip()
        if body:
            yield ExtractedSegment(text=body, page=slide_number, section=title or f"Slide {slide_number}")


def _iter_text(path: Path) -> Iterator[ExtractedSegment]:
    section: str | None = None
    buffer: list[str] = []
    size = 0
    with path.open("r", encoding="utf-8", errors="replace") as handle:
        for line in handle:
            stripped = line.strip()
            if stripped.startswith("#"):
                section = stripped.lstrip("#").strip() or section
            buffer.append(line)
            size += len(line)
            if size >= 64 * 1024:
                text = "".join(buffer).strip()
                if text:
                    yield ExtractedSegment(text=text, section=section)
                buffer = []
                size = 0
    text = "".join(buffer).strip()
    if text:
        yield ExtractedSegment(text=text, section=section)


def _iter_zip(path: Path) -> Iterator[ExtractedSegment]:
    supported = (".txt", ".md", ".markdown", ".csv")
    total_uncompressed = 0
    with zipfile.ZipFile(path) as archive:
        for info in archive.infolist():
            if info.is_dir() or not info.filename.lower().endswith(supported):
                continue
            total_uncompressed += info.file_size
            if total_uncompressed > 200 * 1024 * 1024:
                raise ValueError("Archive expands beyond the 200 MB safety limit")
            with archive.open(info) as member:
                wrapper = io.TextIOWrapper(member, encoding="utf-8", errors="replace")
                buffer: list[str] = []
                size = 0
                for line in wrapper:
                    buffer.append(line)
                    size += len(line)
                    if size >= 64 * 1024:
                        text = "".join(buffer).strip()
                        if text:
                            yield ExtractedSegment(text=text, section=info.filename)
                        buffer = []
                        size = 0
                text = "".join(buffer).strip()
                if text:
                    yield ExtractedSegment(text=text, section=info.filename)


def extract_text_from_file(file_bytes: bytes, file_type: str) -> str:
    """Compatibility helper for callers that still provide in-memory bytes."""
    suffix = {"pdf": ".pdf", "docx": ".docx", "pptx": ".pptx", "zip": ".zip"}.get(normalize_file_type(file_type), ".txt")
    import tempfile

    path = ""
    try:
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as handle:
            handle.write(file_bytes)
            path = handle.name
        return "\n".join(segment.text for segment in iter_extracted_segments(path, file_type)).strip()
    finally:
        if path:
            try:
                os.unlink(path)
            except FileNotFoundError:
                pass


def _clean_metadata_value(value: Any) -> str | None:
    text = str(value).strip() if value is not None else ""
    return text or None
